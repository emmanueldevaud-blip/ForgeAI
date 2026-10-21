"""Service d'import agenda depuis Excel (calendrier des présences)."""
import re
import uuid
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.agenda import AgendaImportBatch, AgendaImportStagedRow, AgendaPresence
from app.models.buildings import Room
from app.models.user import User


MONTHS = {
    "janvier": 1, "fevrier": 2, "février": 2, "mars": 3, "avril": 4, "mai": 5,
    "juin": 6, "juillet": 7, "aout": 8, "août": 8, "septembre": 9,
    "octobre": 10, "novembre": 11, "decembre": 12, "décembre": 12,
}

SITE_NAMES = ("Cébazat", "Riom", "Puteaux")
DAY_NAMES = ("lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche")

GENERIC_PREFIXES = (
    "invité", "invités", "ménage", "menage", "rldc", "intendance",
    "formation", "intervention", "maintenance", "total"
)

CODE_PATTERN = re.compile(r"^(P|D)(-am|-pm)?$", re.I)


def norm(value: Any) -> str:
    if value is None:
        return ""
    value = unicodedata.normalize("NFKD", str(value))
    value = "".join(c for c in value if not unicodedata.combining(c))
    value = re.sub(r"[^a-z0-9]+", " ", value.lower())
    return " ".join(value.split())


def split_name(label: str) -> tuple[str, str]:
    label = " ".join(str(label).split())
    if "," in label:
        last, _, first = label.partition(",")
        return last.strip(), first.strip()
    parts = label.split(" ")
    if len(parts) >= 2:
        return parts[0], " ".join(parts[1:])
    return label, ""


def parse_header_date(value: Any) -> tuple[int, int] | None:
    if isinstance(value, datetime):
        return value.month, value.day
    if isinstance(value, str):
        m = re.match(r"^\s*(\d{1,2})\s+([A-Za-zÀ-ÿ]+)\s*$", value.strip())
        if m and m.group(2).lower() in MONTHS:
            return MONTHS[m.group(2).lower()], int(m.group(1))
    return None


def infer_year(md: tuple[int, int], prev: tuple[int, int, int] | None) -> int:
    """Year of a (month, day) header consistent with previous sheet."""
    month, day = md
    candidates = []
    for y in (2024, 2025, 2026, 2027):
        try:
            d = date(y, month, day)
        except ValueError:
            continue
        if d.weekday() == 0:
            candidates.append(y)
    if not candidates:
        return prev[2] if prev else 2026
    if prev:
        prev_month, _, prev_year = prev
        ordered = [y for y in candidates if y > prev_year or (y == prev_year and month >= prev_month)]
        if ordered:
            return ordered[0]
        return min(candidates, key=lambda y: abs(y - prev_year))
    return candidates[0]


def is_format_b(row2: list, row3: list) -> bool:
    """Format B: day names in row2, site names Cébazat/Riom/Puteaux in row3."""
    days = sum(
        1 for v in row2
        if isinstance(v, str) and norm(v) in DAY_NAMES
    )
    sites = sum(
        1 for v in row3
        if isinstance(v, str) and v.strip() in SITE_NAMES
    )
    return days >= 3 and sites >= 3


def extract_layout(row2: list, row3: list) -> tuple[list[int], dict[int, str]]:
    """Return (day_columns, site_column -> site_name)."""
    day_cols = [
        i + 1 for i, v in enumerate(row2)
        if isinstance(v, str) and norm(v) in DAY_NAMES
    ]
    site_map = {
        i + 1: str(v).strip()
        for i, v in enumerate(row3)
        if isinstance(v, str) and v.strip() in SITE_NAMES
    }
    return day_cols, site_map


def is_generic_label(label: str) -> bool:
    l = norm(label)
    return any(l.startswith(p) for p in GENERIC_PREFIXES)


async def find_user_by_name(db: AsyncSession, label: str) -> list[User]:
    """Find users matching the label. Returns list of candidates."""
    last, first = split_name(label)
    nlast, nfirst = norm(last), norm(first)

    # 1. exact last+first or first+last
    stmt = select(User).where(User.is_active.is_(True))
    users = (await db.execute(stmt)).scalars().all()

    candidates = []
    for u in users:
        u_last, u_first = norm(u.last_name or ""), norm(u.first_name or "")
        if u_last == nlast and u_first == nfirst:
            candidates.append(u)
        elif u_last == nfirst and u_first == nlast:
            candidates.append(u)
    if len(candidates) == 1:
        return candidates

    # 2. first name startswith + last name suffix match
    if not candidates and nfirst:
        for u in users:
            u_last, u_first = norm(u.last_name or ""), norm(u.first_name or "")
            if u_first and (nfirst == u_first or nfirst.startswith(u_first)) and u_last and nlast.endswith(u_last):
                candidates.append(u)
        if len(candidates) == 1:
            return candidates

    # 3. all user tokens ⊆ label tokens
    if not candidates:
        label_tokens = set(norm(label).split())
        for u in users:
            u_tokens = set(norm(f"{u.last_name or ''} {u.first_name or ''}").split())
            if u_tokens and u_tokens <= label_tokens:
                candidates.append(u)

    return candidates


async def find_close_users(db: AsyncSession, label: str, limit: int = 3) -> list[User]:
    """Find close matches for manual matching."""
    from difflib import get_close_matches

    users = (await db.execute(
        select(User).where(User.is_active.is_(True))
    )).scalars().all()
    user_labels = [f"{u.last_name or ''} {u.first_name or ''}".strip() for u in users]
    close = get_close_matches(label, user_labels, n=limit, cutoff=0.72)
    if not close:
        return []
    result = []
    for cl in close:
        for u in users:
            if f"{u.last_name or ''} {u.first_name or ''}".strip() == cl:
                result.append(u)
                break
    return result


async def resolve_site_rooms(db: AsyncSession) -> dict[str, int]:
    """Return {site_label_lower: room_id} for Cébazat/Riom/Puteaux (bureau rooms only)."""
    from sqlalchemy import or_
    from sqlalchemy.orm import selectinload
    from app.models.buildings import Building, UsageType
    BUREAU_USAGE_CODES = ("BUREAUX", "BUREAU")
    stmt = (
        select(Room, Building.name)
        .join(Building, Building.id == Room.building_id)
        .outerjoin(UsageType, UsageType.id == Room.usage_type_id)
        .where(
            Room.is_active.is_(True),
            Building.is_active.is_(True),
            or_(
                UsageType.code.in_(BUREAU_USAGE_CODES),
                UsageType.name.ilike("%bureau%"),
            ),
        )
        .options(selectinload(Room.building))
    )
    rows = (await db.execute(stmt)).all()
    mapping = {}
    for r, building_name in rows:
        bn = (building_name or "").lower()
        room_name = r.name.lower()
        for site in SITE_NAMES:
            nsite = norm(site)
            if nsite in bn or nsite in room_name:
                mapping[nsite] = r.id
    return mapping


async def analyze_import(db: AsyncSession, file_content: bytes, filename: str, uploaded_by: int) -> AgendaImportBatch:
    """Parse the Excel file, stage rows, return batch with bilan."""
    workbook = openpyxl.load_workbook(BytesIO(file_content), data_only=True)

    # Preload users and site rooms
    users = (await db.execute(
        select(User).where(User.is_active.is_(True))
    )).scalars().all()
    site_rooms = await resolve_site_rooms(db)

    batch = AgendaImportBatch(
        reference=uuid.uuid4().hex,
        file_name=filename,
        status="analyzed",
        uploaded_by=uploaded_by,
    )
    db.add(batch)
    await db.flush()

    rows_total = 0
    rows_ready = 0
    rows_needs_review = 0
    rows_skipped_existing = 0
    rows_excluded = 0
    week_dates: list[date] = []

    prev_header: tuple[int, int, int] | None = None

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        if ws.max_row < 6 or ws.max_column < 10:
            continue

        # Read header rows (0-indexed)
        raw_rows = list(ws.iter_rows(values_only=True))
        row1 = list(raw_rows[0]) + [None] * 30
        row2 = list(raw_rows[1]) + [None] * 30
        row3 = list(raw_rows[2]) + [None] * 30

        if not is_format_b(row2, row3):
            continue  # skip format A / empty

        day_cols, site_map = extract_layout(row2, row3)
        if len(day_cols) < 5:
            continue

        # Week dates
        start_md = parse_header_date(row1[1])
        year = prev_header[2] if prev_header else 2026
        if start_md:
            year = infer_year(start_md, prev_header)
            prev_header = (start_md[0], 0, year)
        week_start = date(year, *start_md) if start_md else None
        if week_start and week_start.weekday() != 0:
            continue  # invalid week
        week_days = [date.fromordinal(week_start.toordinal() + i) for i in range(5)] if week_start else []
        if week_start:
            week_dates.append(week_start)
            week_dates.append(week_days[-1])

        # Person rows start at row 4 (0-indexed = 3) for format B
        for r_idx in range(3, len(raw_rows)):
            row = raw_rows[r_idx]
            label = row[0] if row else None
            if label is None or not str(label).strip():
                continue
            label = " ".join(str(label).split())

            # Skip aggregate rows
            if re.match(r"^(total\b|nb de places|semaine$)", label, re.I):
                continue

            # Read values for this row
            values: dict[int, str] = {}
            for dc in day_cols:
                v = row[dc - 1] if dc - 1 < len(row) else None
                if v is not None and str(v).strip():
                    values[dc] = str(v).strip()
                # Site columns: day_col + 1 .. day_col + 3
                for sc in range(dc + 1, dc + 4):
                    if sc in site_map:
                        sv = row[sc - 1] if sc - 1 < len(row) else None
                        if sv is not None and str(sv).strip():
                            values[sc] = str(sv).strip()

            if not values:
                continue


            for d_idx, dc in enumerate(day_cols):
                if d_idx >= len(week_days):
                    break
                presence_date = week_days[d_idx]

                day_code = values.get(dc)
                site_marked = [
                    (sc, site_map[sc])
                    for sc in range(dc + 1, dc + 4)
                    if sc in site_map and values.get(sc)
                ]
                marked_sites = [s for _, s in site_marked]

                # Determine presence type
                is_present = True
                is_remote = False
                period = "full"
                needs_workstation = False
                needs_meal = False
                room_id = None
                status = "ready"
                review_reason = None
                suggestions = None
                user_id = None
                external_name = None

                if day_code and day_code.upper().startswith("P"):
                    is_present = True
                    if day_code.upper().endswith("-AM"):
                        period = "morning"
                    elif day_code.upper().endswith("-PM"):
                        period = "afternoon"
                    needs_workstation = True
                    needs_meal = bool(marked_sites)

                    if not marked_sites:
                        status = "needs_review"
                        review_reason = "site_manquant"
                    elif len(marked_sites) > 1:
                        room_id = site_rooms.get(norm(marked_sites[0]))
                    else:
                        room_id = site_rooms.get(norm(marked_sites[0]))

                elif day_code and day_code.upper().startswith("D"):
                    is_present = True
                    is_remote = True
                    if day_code.upper().endswith("-AM"):
                        period = "morning"
                    elif day_code.upper().endswith("-PM"):
                        period = "afternoon"
                    needs_workstation = False
                    needs_meal = False
                    room_id = None

                elif site_marked and not day_code:
                    # X without P → treat as present
                    is_present = True
                    is_remote = False
                    period = "full"
                    needs_workstation = True
                    needs_meal = True
                    if len(marked_sites) > 1:
                        room_id = site_rooms.get(norm(marked_sites[0]))
                    else:
                        room_id = site_rooms.get(norm(marked_sites[0]))
                    if marked_sites:
                        # flag that day code was missing
                        pass

                elif day_code:
                    # Unknown code - skip
                    continue
                else:
                    # No code, no site mark - skip
                    continue

                # Match user / external
                if is_generic_label(label):
                    external_name = label
                    user_id = None
                else:
                    candidates = await find_user_by_name(db, label)
                    if len(candidates) == 1:
                        user_id = candidates[0].id
                        external_name = None
                    elif len(candidates) > 1:
                        status = "needs_review"
                        review_reason = "utilisateur_ambiguous"
                        suggestions = [{"user_id": c.id, "label": f"{c.last_name} {c.first_name}"} for c in candidates]
                    else:
                        # No exact match - check close matches
                        close = await find_close_users(db, label)
                        if close:
                            status = "needs_review"
                            review_reason = "utilisateur_proche"
                            suggestions = [{"user_id": c.id, "label": f"{c.last_name} {c.first_name}"} for c in close]
                        else:
                            external_name = label
                            user_id = None

                rows_total += 1

                # Check existing in DB
                if status == "ready" and room_id is not None:
                    existing = await db.execute(
                        select(AgendaPresence.id).where(
                            AgendaPresence.presence_date == presence_date,
                            AgendaPresence.room_id == room_id,
                            AgendaPresence.user_id == (user_id if user_id else -1)
                        )
                    )
                    # For external: check external_name match (user_id IS NULL)
                    if user_id is None and external_name:
                        existing = await db.execute(
                            select(AgendaPresence.id).where(
                                AgendaPresence.presence_date == presence_date,
                                AgendaPresence.room_id == room_id,
                                AgendaPresence.user_id.is_(None),
                                AgendaPresence.external_name == external_name
                            )
                        )
                    elif user_id:
                        existing = await db.execute(
                            select(AgendaPresence.id).where(
                                AgendaPresence.presence_date == presence_date,
                                AgendaPresence.room_id == room_id,
                                AgendaPresence.user_id == user_id
                            )
                        )
                    else:
                        existing = None
                    if existing and existing.scalar_one_or_none():
                        status = "skipped_existing"

                # Update counters
                if status == "ready":
                    rows_ready += 1
                elif status == "needs_review":
                    rows_needs_review += 1
                elif status == "skipped_existing":
                    rows_skipped_existing += 1

                staged = AgendaImportStagedRow(
                    batch_id=batch.id,
                    sheet_name=sheet_name,
                    line_number=r_idx + 1,
                    presence_date=presence_date,
                    period=period,
                    is_remote=is_remote,
                    raw_label=label,
                    user_id=user_id,
                    external_name=external_name,
                    room_id=room_id,
                    site_label=site_marked[0][1] if site_marked else None,
                    needs_workstation=needs_workstation,
                    needs_meal=needs_meal,
                    status=status,
                    review_reason=review_reason,
                    suggestions=suggestions,
                )
                db.add(staged)

    # Counters were accumulated during parsing
    batch.rows_total = rows_total
    batch.rows_ready = rows_ready
    batch.rows_needs_review = rows_needs_review
    batch.rows_skipped_existing = rows_skipped_existing
    batch.rows_excluded = rows_excluded
    if week_dates:
        batch.week_start = min(week_dates)
        batch.week_end = max(week_dates)

    await db.commit()
    await db.refresh(batch)
    return batch


async def get_batch(db: AsyncSession, batch_id: int) -> AgendaImportBatch | None:
    return await db.get(AgendaImportBatch, batch_id)


async def get_staged_rows(
    db: AsyncSession,
    batch_id: int,
    status: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    limit: int = 200,
    offset: int = 0,
) -> list[AgendaImportStagedRow]:
    stmt = select(AgendaImportStagedRow).where(AgendaImportStagedRow.batch_id == batch_id)
    if status:
        stmt = stmt.where(AgendaImportStagedRow.status == status)
    if start_date:
        stmt = stmt.where(AgendaImportStagedRow.presence_date >= start_date)
    if end_date:
        stmt = stmt.where(AgendaImportStagedRow.presence_date <= end_date)
    stmt = stmt.order_by(AgendaImportStagedRow.presence_date, AgendaImportStagedRow.sheet_name, AgendaImportStagedRow.line_number)
    stmt = stmt.limit(limit).offset(offset)
    return (await db.execute(stmt)).scalars().all()


async def update_staged_row(
    db: AsyncSession,
    row_id: int,
    user_id: int | None = None,
    external_name: str | None = None,
    room_id: int | None = None,
    status: str | None = None,
    skip: bool = False,
) -> AgendaImportStagedRow | None:
    row = await db.get(AgendaImportStagedRow, row_id)
    if not row:
        return None
    if user_id is not None:
        row.user_id = user_id
        row.external_name = None
    if external_name is not None:
        row.external_name = external_name
        row.user_id = None
    if room_id is not None:
        row.room_id = room_id
    if status is not None:
        row.status = status
    if skip:
        row.status = "excluded"
    # If resolved, mark ready
    if row.status == "needs_review" and row.user_id is not None and row.room_id is not None:
        row.status = "ready"
    await db.commit()
    await db.refresh(row)
    return row


async def validate_import(db: AsyncSession, batch_id: int, current_user_id: int) -> tuple[int, int]:
    """Create AgendaPresence for ready rows. Returns (created, skipped)."""
    batch = await db.get(AgendaImportBatch, batch_id)
    if not batch or batch.status == "validated":
        return (0, 0)

    rows = await db.execute(
        select(AgendaImportStagedRow).where(
            AgendaImportStagedRow.batch_id == batch_id,
            AgendaImportStagedRow.status == "ready"
        )
    )
    staged_rows = rows.scalars().all()

    created = 0
    skipped = 0
    for row in staged_rows:
        # Double-check existing
        existing = await db.execute(
            select(AgendaPresence.id).where(
                AgendaPresence.presence_date == row.presence_date,
                AgendaPresence.room_id == row.room_id,
                AgendaPresence.user_id == (row.user_id if row.user_id else -1)
            )
        )
        if row.user_id is None and row.external_name:
            existing = await db.execute(
                select(AgendaPresence.id).where(
                    AgendaPresence.presence_date == row.presence_date,
                    AgendaPresence.room_id == row.room_id,
                    AgendaPresence.user_id.is_(None),
                    AgendaPresence.external_name == row.external_name
                )
            )
        if existing.scalar_one_or_none():
            row.status = "skipped_existing"
            skipped += 1
            continue

        presence = AgendaPresence(
            presence_date=row.presence_date,
            room_id=row.room_id,
            user_id=row.user_id,
            external_name=row.external_name,
            source="manual",
            source_ref=batch.reference,
            is_present=True,
            needs_workstation=row.needs_workstation,
            needs_meal=row.needs_meal,
            is_remote=row.is_remote,
            period=row.period,
            created_by=current_user_id,
        )
        db.add(presence)
        await db.flush()
        row.status = "validated"
        row.created_presence_id = presence.id
        created += 1

    batch.status = "validated"
    batch.validated_at = datetime.now(timezone.utc)
    batch.created_count = created
    await db.commit()
    return (created, skipped)