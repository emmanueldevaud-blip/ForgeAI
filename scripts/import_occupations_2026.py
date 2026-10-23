#!/usr/bin/env python3
"""
Import occupations 2026 from Excel to Production.
Run on VPS: docker cp ... && docker exec ... python /app/scripts/import_occupations_2026.py
"""
import asyncio
from datetime import date, datetime, timedelta
from collections import defaultdict
import openpyxl
import re

from app.db.session import AsyncSessionLocal
from app.models.housing import Housing, Occupant, Occupancy, occupancy_occupants
from app.models.buildings import Room
from sqlalchemy import select, and_, or_, delete
from sqlalchemy.orm import selectinload
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.dialects.mysql import insert as mysql_insert


# ============================================================
# CONFIGURATION
# ============================================================

# Room name mapping: Excel column -> (Housing name, room_index)
ROOM_MAPPING = {
    "Les Sentiers  - Chambre 1": ("LES SENTIERS", 0),
    "Les Sentiers  - Chambre 2": ("LES SENTIERS", 1),
    "Le Puy de Dôme ": ("LE PUY DE DÔME ", 0),
    "Les Sources - Chambre 1": ("LES SOURCES", 0),
    "Les Sources - Chambre 2": ("LES SOURCES", 1),
    "Le Pariou - Chambre 2": ("LE PARIOU", 1),
    "Le Pariou - Chambre 1": ("LE PARIOU", 0),
    " Le Sancy  - Chambre 1": ("LE SANCY", 0),
}

# Occupants to skip
SKIP_OCCUPANTS = {"Préparation chambre", "Camille POPULLO"}

# Names to split (couples/families) - format: "Name in Excel" -> [(first, last), ...]
OCCUPANT_SPLITS = {
    "Asaph et Cécile YEPMOU": [("Asaph", "YEPMOU"), ("Cécile", "YEPMOU")],
    "Famille STHOREZ": [("Famille", "STHOREZ")],  # single entry, no first names
    "Famille VANDEPOORTE": [("Famille", "VANDEPOORTE")],
    "Gaëlle et Gwendal TILLY": [("Gaëlle", "TILLY"), ("Gwendal", "TILLY")],
    "Jonathan et Amandine ROUZIER": [("Jonathan", "ROUZIER"), ("Amandine", "ROUZIER")],
    "Joffrey et Isa ROUX": [("Joffrey", "ROUX"), ("Isa", "ROUX")],
    "Julien et Anna WINE": [("Julien", "WINE"), ("Anna", "WINE")],
    "Julien et Jessy DUMITRAN": [("Julien", "DUMITRAN"), ("Jessy", "DUMITRAN")],
    "Lydie et Sébastien MASSIAUX": [("Lydie", "MASSIAUX"), ("Sébastien", "MASSIAUX")],
    "Nat et Caro DESMERY": [("Nat", "DESMERY"), ("Caro", "DESMERY")],
    "Oudish APPADOO": [("Oudish", "APPADOO")],
    "Philippe et Manuela RETAILLEAU": [("Philippe", "RETAILLEAU"), ("Manuela", "RETAILLEAU")],
    "Robin DHURMAJEE": [("Robin", "DHURMAJEE")],
    "Yann et Aurélie HOCHENAUER": [("Yann", "HOCHENAUER"), ("Aurélie", "HOCHENAUER")],
    "Arthur GOUVERNELLE": [("Arthur", "GOUVERNELLE")],
    "Yann HOCHENAUER": [("Yann", "HOCHENAUER")],
    "Théo FIZZAROTTI": [("Théo", "FIZZAROTTI")],
    "Maxence FOUQUET": [("Maxence", "FOUQUET")],
    "Manon BONDE": [("Manon", "BONDE")],
    "Allan YINDOULA": [("Allan", "YINDOULA")],
    "Erwan LUCY": [("Erwan", "LUCY")],
    "Wendy YINDOULA": [("Wendy", "YINDOULA")],
    "BUDIN": [("BUDIN", "")],  # no last name
    "BUDIN Frederic": [("Frederic", "BUDIN")],
    "Franck MICHEL": [("Franck", "MICHEL")],
    "Katia MALDUE": [("Katia", "MALDUE")],
    "Mathéo CARNEIRO": [("Mathéo", "CARNEIRO")],
    "Sylvain THOMAS": [("Sylvain", "THOMAS")],
    "Annie BARRIERE": [("Annie", "BARRIERE")],
    "DESEUCHE": [("DESEUCHE", "")],
    "Eric MASSON": [("Eric", "MASSON")],
    "Estelle HOAREAU": [("Estelle", "HOAREAU")],
    "Marion VO": [("Marion", "VO")],
    "Corentin JAGER": [("Corentin", "JAGER")],
    "Loïc et Cynthia PARTEL": [("Loïc", "PARTEL"), ("Cynthia", "PARTEL")],
    "Yannis LUCY": [("Yannis", "LUCY")],
}


def split_occupant_name(excel_name):
    """Split Excel occupant name into list of (first_name, last_name) tuples."""
    name = excel_name.strip()
    if name in OCCUPANT_SPLITS:
        return OCCUPANT_SPLITS[name]
    # Default: try to parse "First Last"
    parts = name.split()
    if len(parts) >= 2:
        return [(" ".join(parts[:-1]), parts[-1])]
    return [(name, "")]


async def get_or_create_occupant(db, first_name, last_name):
    """Get existing occupant or create new one."""
    result = await db.execute(
        select(Occupant).where(
            Occupant.first_name == first_name,
            Occupant.last_name == last_name
        )
    )
    occupant = result.scalar_one_or_none()
    if occupant:
        return occupant
    occupant = Occupant(first_name=first_name, last_name=last_name, is_active=True)
    db.add(occupant)
    await db.flush()
    return occupant


async def get_housing_by_room_name(db, room_excel_name):
    """Get housing_id and room_index from Excel room name."""
    mapping = ROOM_MAPPING.get(room_excel_name)
    if not mapping:
        return None, None
    housing_name, room_index = mapping
    result = await db.execute(
        select(Housing).join(Room).where(Room.name == housing_name)
    )
    housing = result.scalar_one_or_none()
    if housing:
        return housing.id, room_index
    return None, None


async def main():
    # Load Excel
    wb = openpyxl.load_workbook('/app/Occupation_Logements_2026_corrige.xlsx', read_only=True, data_only=True)
    ws = wb.active
    
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    
    print(f"Loaded {len(rows)} rows, {len(headers)-1} rooms")
    
    # Parse Excel: build stays per (room, occupant)
    # Structure: stays[(room, occupant)] = list of (date, row_idx)
    stays_raw = defaultdict(list)
    
    for row_idx, row in enumerate(rows):
        dt = row[0]
        if not dt:
            continue
        d = dt.date() if isinstance(dt, datetime) else dt
        for col_idx in range(1, len(headers)):
            val = row[col_idx]
            if not val or not str(val).strip():
                continue
            name = str(val).strip()
            if name in SKIP_OCCUPANTS:
                continue
            room_name = headers[col_idx]
            stays_raw[(room_name, name)].append((d, row_idx + 2))
    
    print(f"Raw stays found: {len(stays_raw)}")
    
    # Group consecutive days into stays
    stays = []  # list of (room_name, occupant_name, arrival, departure)
    for (room, occ), days in stays_raw.items():
        days_sorted = sorted(days, key=lambda x: x[0])
        stay_start = days_sorted[0][0]
        prev_day = days_sorted[0][0]
        for d, _ in days_sorted[1:]:
            if (d - prev_day).days > 1:
                # Gap - end current stay
                stays.append((room, occ, stay_start, prev_day))
                stay_start = d
            prev_day = d
        # Last stay
        stays.append((room, occ, stay_start, prev_day))
    
    print(f"Grouped stays: {len(stays)}")
    
    async with AsyncSessionLocal() as db:
        # Load existing housing mapping
        housing_cache = {}
        for room_excel, (housing_name, room_idx) in ROOM_MAPPING.items():
            hid, ridx = await get_housing_by_room_name(db, room_excel)
            if hid:
                housing_cache[room_excel] = (hid, ridx)
                print(f"  {room_excel} -> Housing ID={hid}, room_index={ridx}")
            else:
                print(f"  WARNING: {room_excel} -> Housing NOT FOUND")
        
        # Load existing occupants
        existing_occupants = {}
        result = await db.execute(select(Occupant))
        for occ in result.scalars().all():
            key = f"{occ.first_name} {occ.last_name}".strip()
            existing_occupants[key.lower()] = occ
        
        # Handle Yannis LUCY merge: find existing occupancy
        yannis_occ = None
        result = await db.execute(
            select(Occupancy).join(Housing).join(Room)
            .options(selectinload(Occupancy.occupants))
            .where(Room.name == "LES SENTIERS")
            .join(Occupancy.occupants).where(Occupant.first_name == "Yannis", Occupant.last_name == "LUCY")
        )
        yannis_occ = result.scalar_one_or_none()
        if yannis_occ:
            print(f"Found existing Yannis LUCY occupancy: ID={yannis_occ.id}, {yannis_occ.arrival_date.date()} to {yannis_occ.departure_date.date()}")
        
        # Process stays
        created_occupants = 0
        created_occupancies = 0
        linked = 0
        skipped = 0
        
        for room_excel, occ_name, arrival, departure in stays:
            if room_excel not in housing_cache:
                print(f"  SKIP: No housing for {room_excel}")
                skipped += 1
                continue
            
            housing_id, room_index = housing_cache[room_excel]
            
            # Split occupant name
            occ_parts = split_occupant_name(occ_name)
            occupant_objs = []
            for first, last in occ_parts:
                key = f"{first} {last}".strip().lower()
                if key in existing_occupants:
                    occ_obj = existing_occupants[key]
                else:
                    occ_obj = await get_or_create_occupant(db, first, last)
                    existing_occupants[key] = occ_obj
                    created_occupants += 1
                    print(f"  Created occupant: {first} {last}")
                occupant_objs.append(occ_obj)
            
            # Special handling: Yannis LUCY in Les Sentiers Chambre 1
            if occ_name == "Yannis LUCY" and room_excel == "Les Sentiers  - Chambre 1" and yannis_occ:
                # Extend existing occupancy to cover full year
                if arrival < yannis_occ.arrival_date.date():
                    yannis_occ.arrival_date = datetime.combine(arrival, datetime.min.time())
                if departure > yannis_occ.departure_date.date():
                    yannis_occ.departure_date = datetime.combine(departure, datetime.max.time())
                yannis_occ.room_index = 0
                yannis_occ.status = "confirmed"
                yannis_occ.nb_persons = len(occupant_objs)
                # Ensure link exists
                for occ_obj in occupant_objs:
                    if occ_obj not in yannis_occ.occupants:
                        yannis_occ.occupants.append(occ_obj)
                        linked += 1
                print(f"  Merged/Extended Yannis LUCY occupancy {yannis_occ.id}: {arrival} to {departure}")
                continue
            
            # Check for existing overlapping occupancy
            result = await db.execute(
                select(Occupancy)
                .options(selectinload(Occupancy.occupants))
                .where(
                    Occupancy.housing_id == housing_id,
                    Occupancy.room_index == room_index,
                    Occupancy.arrival_date <= datetime.combine(departure, datetime.max.time()),
                    Occupancy.departure_date >= datetime.combine(arrival, datetime.min.time()),
                    Occupancy.status.in_(["confirmed", "pre_reserved", "in_progress"])
                )
            )
            existing = result.scalar_one_or_none()
            
            if existing:
                # Extend existing
                if arrival < existing.arrival_date.date():
                    existing.arrival_date = datetime.combine(arrival, datetime.min.time())
                if departure > existing.departure_date.date():
                    existing.departure_date = datetime.combine(departure, datetime.max.time())
                existing.status = "confirmed"
                existing.nb_persons = max(existing.nb_persons, len(occupant_objs))
                for occ_obj in occupant_objs:
                    if occ_obj not in existing.occupants:
                        existing.occupants.append(occ_obj)
                        linked += 1
                print(f"  Extended occupancy {existing.id}: {arrival} to {departure}")
            else:
                # Create new occupancy
                occ = Occupancy(
                    housing_id=housing_id,
                    room_index=room_index,
                    arrival_date=datetime.combine(arrival, datetime.min.time()),
                    departure_date=datetime.combine(departure, datetime.max.time()),
                    status="confirmed",
                    nb_persons=len(occupant_objs),
                    occupants=occupant_objs
                )
                db.add(occ)
                created_occupancies += 1
                linked += len(occupant_objs)
                print(f"  Created occupancy: {room_excel} / {occ_name} ({arrival} to {departure})")
        
        await db.commit()
        print(f"\n=== IMPORT SUMMARY ===")
        print(f"Created occupants: {created_occupants}")
        print(f"Created occupancies: {created_occupancies}")
        print(f"Linked occupants: {linked}")
        print(f"Skipped: {skipped}")

if __name__ == "__main__":
    asyncio.run(main())