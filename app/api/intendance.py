"""API du suivi des dépenses de l'Intendance."""

from datetime import date, timedelta

import math
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from sqlalchemy import String, cast, func, or_, select, update as sa_update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.deps import require_permission
from app.db.session import get_db
from app.models import (
    IntendanceAccountCode, IntendanceCardAlias, IntendanceExpense,
    IntendanceImportBatch, IntendanceImportStagedRow, IntendanceSite, IntendanceCategory, IntendanceDashboardSetting, User,
)
from app.schemas.intendance_expense import (
    AccountCodeCreate, AccountCodeListResponse, AccountCodeResponse, AccountCodeUpdate,
    BulkSiteAssign, BulkSiteResponse, ExpenseListResponse, ExpenseResponse,
    ExpenseUpdate, ImportBatchResponse, ImportValidateResponse, SiteCreate,
    SiteListResponse, SiteResponse, SiteUpdate, StagedRowListResponse, StagedRowResponse, ImportBilan,
    CardAliasCreate, CardAliasListResponse, CardAliasResponse, CardAliasUpdate,
)
from app.services.intendance_expense import expense_duplicate_key, expense_fingerprint, normalize_date, normalize_text, parse_excel_row

router = APIRouter(prefix="/intendance/expenses", tags=["intendance-expenses"])


def _pages(total: int, page_size: int) -> int:
    return math.ceil(total / page_size) if total else 1


def _expense_response(expense: IntendanceExpense) -> ExpenseResponse:
    missing = []
    if not expense.category:
        missing.append("category")
    if not expense.user_name:
        missing.append("user")
    if not expense.site_id:
        missing.append("site")
    return ExpenseResponse.model_validate(expense).model_copy(update={
        "import_reference": expense.import_batch.reference if expense.import_batch else None,
        "missing_fields": missing,
        "is_complete": not missing,
    })


@router.get("", response_model=ExpenseListResponse)
async def list_expenses(
    page: int = 1, page_size: int = 50, search: str | None = None,
    missing: str | None = None, site_id: int | None = None,
    expense_date_from: date | None = None, expense_date_to: date | None = None,
    sort_by: str = "expense_date", sort_order: str = "desc",
    current_user: User = Depends(require_permission("intendance.expenses.view")),
    db: AsyncSession = Depends(get_db),
):
    page = max(page, 1); page_size = min(max(page_size, 1), 500)
    query = select(IntendanceExpense).options(selectinload(IntendanceExpense.site), selectinload(IntendanceExpense.import_batch))
    if search:
        term = f"%{search.strip()}%"
        query = query.where(
            (IntendanceExpense.item_description.ilike(term))
            | (cast(IntendanceExpense.req_number, String).ilike(term))
            | (cast(IntendanceExpense.amount, String).ilike(term))
        )
    if site_id is not None:
        query = query.where(IntendanceExpense.site_id == site_id)
    if missing == "site": query = query.where(IntendanceExpense.site_id.is_(None))
    if missing == "category": query = query.where(IntendanceExpense.category.is_(None))
    if missing == "user": query = query.where(IntendanceExpense.user_name.is_(None))
    if expense_date_from is not None:
        query = query.where(IntendanceExpense.expense_date >= expense_date_from)
    if expense_date_to is not None:
        query = query.where(IntendanceExpense.expense_date <= expense_date_to)
    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()
    sort_columns = {
        "expense_date": IntendanceExpense.expense_date,
        "req_number": IntendanceExpense.req_number,
        "item_description": IntendanceExpense.item_description,
        "amount": IntendanceExpense.amount,
        "account_code": IntendanceExpense.account_code,
        "category": IntendanceExpense.category,
        "user_name": IntendanceExpense.user_name,
        "id": IntendanceExpense.id,
    }
    if sort_by == "site.name":
        query = query.outerjoin(IntendanceSite, IntendanceSite.id == IntendanceExpense.site_id)
        sort_column = IntendanceSite.name
    else:
        sort_column = sort_columns.get(sort_by, IntendanceExpense.expense_date)
    order_clause = sort_column.asc() if sort_order == "asc" else sort_column.desc()
    rows = (await db.execute(query.order_by(order_clause, IntendanceExpense.id.desc()).offset((page - 1) * page_size).limit(page_size))).scalars().all()
    return ExpenseListResponse(items=[_expense_response(row) for row in rows], total=total, page=page, page_size=page_size, total_pages=_pages(total, page_size))


@router.get("/dashboard", response_model=dict)
async def expense_dashboard(
    expense_date_from: date | None = None, expense_date_to: date | None = None,
    user_name: list[str] | None = Query(None), site_id: list[int] | None = Query(None),
    category: list[str] | None = Query(None),
    current_user: User = Depends(require_permission("intendance.expenses.view")),
    db: AsyncSession = Depends(get_db),
):
    """Retourne les indicateurs agrégés du tableau de bord des dépenses."""
    excluded = (await db.execute(select(IntendanceCategory.name).where(IntendanceCategory.excluded_from_dashboard.is_(True)))).scalars().all()
    dashboard_setting = await db.get(IntendanceDashboardSetting, 1)
    exclude_uncategorized = dashboard_setting.exclude_uncategorized if dashboard_setting else False
    def filtered(statement, from_date=expense_date_from, to_date=expense_date_to):
        if from_date is not None:
            statement = statement.where(IntendanceExpense.expense_date >= from_date)
        if to_date is not None:
            statement = statement.where(IntendanceExpense.expense_date <= to_date)
        if user_name:
            statement = statement.where(IntendanceExpense.user_name.in_(user_name))
        if site_id:
            statement = statement.where(IntendanceExpense.site_id.in_(site_id))
        if category:
            statement = statement.where(IntendanceExpense.category.in_(category))
        return statement

    visible_category = IntendanceExpense.category.notin_(excluded)
    if not exclude_uncategorized:
        visible_category = or_(IntendanceExpense.category.is_(None), visible_category)
    base = filtered(select(IntendanceExpense).where(visible_category))
    base_rows = base.subquery()
    totals = (await db.execute(select(
        func.coalesce(func.sum(base_rows.c.amount), 0),
        func.count(base_rows.c.id),
        func.coalesce(func.avg(base_rows.c.amount), 0),
    ).select_from(base_rows))).one()
    category_rows = (await db.execute(filtered(select(
        IntendanceExpense.category,
        func.count(IntendanceExpense.id).label("count"),
        func.sum(IntendanceExpense.amount).label("amount"),
        func.avg(IntendanceExpense.amount).label("average"),
    ).where(visible_category).group_by(IntendanceExpense.category)
    ))).all()
    total_amount = float(totals[0] or 0)
    categories = [{"name": row[0] or "Non classée", "count": row[1], "amount": float(row[2] or 0), "average": float(row[3] or 0)} for row in category_rows]
    categories.sort(key=lambda row: row["amount"], reverse=True)

    monthly_rows = (await db.execute(filtered(select(
        func.extract("year", IntendanceExpense.expense_date).label("year"),
        func.extract("month", IntendanceExpense.expense_date).label("month"),
        func.sum(IntendanceExpense.amount).label("amount"),
        func.count(IntendanceExpense.id).label("count"),
    ).group_by("year", "month").order_by("year", "month")))).all()
    monthly = [{"year": int(row[0]), "month": int(row[1]), "amount": float(row[2] or 0), "count": row[3]} for row in monthly_rows]
    if expense_date_from is not None and expense_date_to is not None and expense_date_to >= expense_date_from:
        monthly_by_key = {(item["year"], item["month"]): item for item in monthly}
        cursor = expense_date_from.replace(day=1)
        last = expense_date_to.replace(day=1)
        complete_months = []
        while cursor <= last:
            complete_months.append(monthly_by_key.get((cursor.year, cursor.month), {"year": cursor.year, "month": cursor.month, "amount": 0, "count": 0}))
            cursor = cursor.replace(year=cursor.year + (cursor.month == 12), month=1 if cursor.month == 12 else cursor.month + 1)
        monthly = complete_months

    user_rows = (await db.execute(filtered(select(
        IntendanceExpense.user_name,
        func.sum(IntendanceExpense.amount).label("amount"),
        func.count(IntendanceExpense.id).label("count"),
    ).where(IntendanceExpense.user_name.is_not(None)).group_by(IntendanceExpense.user_name).order_by(func.sum(IntendanceExpense.amount).desc())))).all()
    users = [{"name": row[0], "amount": float(row[1] or 0), "count": row[2]} for row in user_rows]

    top_rows = (await db.execute(filtered(select(
        IntendanceExpense.id, IntendanceExpense.expense_date, IntendanceExpense.item_description,
        IntendanceExpense.category, IntendanceExpense.user_name, IntendanceExpense.amount,
    ).order_by(IntendanceExpense.amount.desc()).limit(10)))).all()
    top_expenses = [{"id": row[0], "date": row[1].isoformat(), "description": row[2], "category": row[3] or "Non classée", "user_name": row[4], "amount": float(row[5] or 0)} for row in top_rows]

    previous_amount = None
    if expense_date_from is not None and expense_date_to is not None and expense_date_to >= expense_date_from:
        span = (expense_date_to - expense_date_from).days + 1
        previous_from = expense_date_from - timedelta(days=span)
        previous_to = expense_date_from - timedelta(days=1)
        previous_amount = float((await db.execute(filtered(select(func.coalesce(func.sum(IntendanceExpense.amount), 0)).where(visible_category), previous_from, previous_to))).scalar_one() or 0)
    return {"kpis": {"total_amount": total_amount, "count": totals[1], "average_amount": float(totals[2] or 0), "previous_amount": previous_amount}, "categories": categories, "monthly": monthly, "users": users, "top_expenses": top_expenses}


@router.patch("/{expense_id}", response_model=ExpenseResponse)
async def update_expense(expense_id: int, data: ExpenseUpdate, current_user: User = Depends(require_permission("intendance.expenses.edit")), db: AsyncSession = Depends(get_db)):
    expense = (await db.execute(select(IntendanceExpense).options(selectinload(IntendanceExpense.site), selectinload(IntendanceExpense.import_batch)).where(IntendanceExpense.id == expense_id))).scalar_one_or_none()
    if not expense: raise HTTPException(404, "Dépense introuvable")
    changes = data.model_dump(exclude_unset=True)
    if "category" in changes: expense.category_is_manual = True
    if "user_name" in changes: expense.user_is_manual = True
    if "site_id" in changes and changes["site_id"] is not None and not await db.get(IntendanceSite, changes["site_id"]): raise HTTPException(404, "Site introuvable")
    for key, value in changes.items(): setattr(expense, key, value)
    expense.updated_by = current_user.id
    await db.commit(); await db.refresh(expense)
    return _expense_response(expense)


@router.delete("/{expense_id}", status_code=204)
async def delete_expense(expense_id: int, current_user: User = Depends(require_permission("intendance.expenses.edit")), db: AsyncSession = Depends(get_db)):
    expense = await db.get(IntendanceExpense, expense_id)
    if not expense: raise HTTPException(404, "Dépense introuvable")
    await db.delete(expense)
    await db.commit()


@router.post("/sites/bulk", response_model=BulkSiteResponse)
async def bulk_site(data: BulkSiteAssign, current_user: User = Depends(require_permission("intendance.expenses.edit")), db: AsyncSession = Depends(get_db)):
    if not await db.get(IntendanceSite, data.site_id): raise HTTPException(404, "Site introuvable")
    result = await db.execute(select(IntendanceExpense).where(IntendanceExpense.id.in_(data.ids)))
    rows = result.scalars().all()
    for row in rows: row.site_id = data.site_id; row.updated_by = current_user.id
    await db.commit()
    return BulkSiteResponse(updated=len(rows))


@router.get("/sites", response_model=SiteListResponse)
async def list_sites(current_user: User = Depends(require_permission("intendance.expenses.view")), db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(IntendanceSite).order_by(IntendanceSite.sort_order, IntendanceSite.name))).scalars().all()
    return SiteListResponse(items=[SiteResponse.model_validate(x) for x in rows], total=len(rows), page=1, page_size=len(rows) or 1, total_pages=1)


@router.post("/sites", response_model=SiteResponse, status_code=201)
async def create_site(data: SiteCreate, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    site = IntendanceSite(**data.model_dump()); db.add(site)
    try: await db.commit()
    except IntegrityError: await db.rollback(); raise HTTPException(409, "Ce site existe déjà")
    await db.refresh(site); return SiteResponse.model_validate(site)


@router.patch("/sites/{site_id}", response_model=SiteResponse)
async def update_site(site_id: int, data: SiteUpdate, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    site = await db.get(IntendanceSite, site_id)
    if not site: raise HTTPException(404, "Site introuvable")
    for key, value in data.model_dump(exclude_unset=True).items(): setattr(site, key, value)
    await db.commit(); await db.refresh(site); return SiteResponse.model_validate(site)


@router.delete("/sites/{site_id}", status_code=204)
async def delete_site(site_id: int, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    site = await db.get(IntendanceSite, site_id)
    if not site: raise HTTPException(404, "Site introuvable")
    site.is_active = False
    await db.commit()


@router.get("/account-codes", response_model=AccountCodeListResponse)
async def list_account_codes(current_user: User = Depends(require_permission("intendance.expenses.view")), db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(IntendanceAccountCode).order_by(IntendanceAccountCode.code))).scalars().all()
    return AccountCodeListResponse(items=[AccountCodeResponse.model_validate(x) for x in rows], total=len(rows), page=1, page_size=len(rows) or 1, total_pages=1)

@router.get("/categories", response_model=list[dict])
async def list_categories(current_user: User = Depends(require_permission("intendance.expenses.view")), db: AsyncSession = Depends(get_db)):
    setting = await db.get(IntendanceDashboardSetting, 1)
    rows = [{"id": row.id, "name": row.name, "excluded_from_dashboard": row.excluded_from_dashboard} for row in (await db.execute(select(IntendanceCategory).order_by(IntendanceCategory.name))).scalars()]
    rows.append({"id": "uncategorized", "name": "Non classée", "excluded_from_dashboard": setting.exclude_uncategorized if setting else False, "system": True})
    return rows

@router.patch("/categories/uncategorized", response_model=dict)
async def update_uncategorized_category(data: dict, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    setting = await db.get(IntendanceDashboardSetting, 1)
    if setting is None:
        setting = IntendanceDashboardSetting(id=1)
        db.add(setting)
    setting.exclude_uncategorized = bool(data.get("excluded_from_dashboard", setting.exclude_uncategorized))
    await db.commit()
    return {"id": "uncategorized", "name": "Non classée", "excluded_from_dashboard": setting.exclude_uncategorized, "system": True}

@router.post("/categories", response_model=dict, status_code=201)
async def create_category(data: dict, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    category = IntendanceCategory(name=str(data.get("name", "")).strip(), excluded_from_dashboard=bool(data.get("excluded_from_dashboard", False)))
    if not category.name: raise HTTPException(422, "Le nom est obligatoire")
    db.add(category)
    try: await db.commit()
    except IntegrityError: await db.rollback(); raise HTTPException(409, "Cette catégorie existe déjà")
    await db.refresh(category); return {"id": category.id, "name": category.name}

@router.patch("/categories/{category_id}", response_model=dict)
async def update_category(category_id: int, data: dict, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    category = await db.get(IntendanceCategory, category_id)
    if not category: raise HTTPException(404, "Catégorie introuvable")
    name = str(data.get("name", category.name)).strip()
    if not name: raise HTTPException(422, "Le nom est obligatoire")
    await db.execute(sa_update(IntendanceAccountCode).where(IntendanceAccountCode.category == category.name).values(category=name))
    await db.execute(sa_update(IntendanceExpense).where(IntendanceExpense.category == category.name).values(category=name))
    category.name = name
    category.excluded_from_dashboard = bool(data.get("excluded_from_dashboard", category.excluded_from_dashboard))
    try: await db.commit()
    except IntegrityError: await db.rollback(); raise HTTPException(409, "Cette catégorie existe déjà")
    return {"id": category.id, "name": category.name}

@router.delete("/categories/{category_id}", status_code=204)
async def delete_category(category_id: int, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    category = await db.get(IntendanceCategory, category_id)
    if not category: raise HTTPException(404, "Catégorie introuvable")
    used = (await db.execute(select(func.count()).select_from(IntendanceAccountCode).where(IntendanceAccountCode.category == category.name))).scalar_one()
    used += (await db.execute(select(func.count()).select_from(IntendanceExpense).where(IntendanceExpense.category == category.name))).scalar_one()
    if used: raise HTTPException(409, "Cette catégorie est utilisée et ne peut pas être supprimée")
    await db.delete(category); await db.commit()


@router.get("/aliases", response_model=CardAliasListResponse)
async def list_aliases(current_user: User = Depends(require_permission("intendance.expenses.view")), db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(select(IntendanceCardAlias).order_by(IntendanceCardAlias.alias))).scalars().all()
    return CardAliasListResponse(items=[CardAliasResponse.model_validate(x) for x in rows], total=len(rows), page=1, page_size=len(rows) or 1, total_pages=1)


@router.post("/aliases", response_model=CardAliasResponse, status_code=201)
async def create_alias(data: CardAliasCreate, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    alias = IntendanceCardAlias(**data.model_dump()); db.add(alias)
    try: await db.commit()
    except IntegrityError: await db.rollback(); raise HTTPException(409, "Cet alias existe déjà")
    await db.refresh(alias); return CardAliasResponse.model_validate(alias)


@router.patch("/aliases/{alias_id}", response_model=CardAliasResponse)
async def update_alias(alias_id: int, data: CardAliasUpdate, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    alias = await db.get(IntendanceCardAlias, alias_id)
    if not alias: raise HTTPException(404, "Alias introuvable")
    for key, value in data.model_dump(exclude_unset=True).items(): setattr(alias, key, value)
    await db.commit(); await db.refresh(alias); return CardAliasResponse.model_validate(alias)


@router.delete("/aliases/{alias_id}", status_code=204)
async def delete_alias(alias_id: int, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    alias = await db.get(IntendanceCardAlias, alias_id)
    if not alias: raise HTTPException(404, "Alias introuvable")
    await db.delete(alias); await db.commit()


@router.post("/account-codes", response_model=AccountCodeResponse, status_code=201)
async def create_account_code(data: AccountCodeCreate, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    code = IntendanceAccountCode(**data.model_dump()); db.add(code)
    try: await db.commit()
    except IntegrityError: await db.rollback(); raise HTTPException(409, "Ce code existe déjà")
    await db.refresh(code); return AccountCodeResponse.model_validate(code)


@router.patch("/account-codes/{code_id}", response_model=AccountCodeResponse)
async def update_account_code(code_id: int, data: AccountCodeUpdate, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    code = await db.get(IntendanceAccountCode, code_id)
    if not code: raise HTTPException(404, "Code comptable introuvable")
    for key, value in data.model_dump(exclude_unset=True).items(): setattr(code, key, value)
    await db.commit(); await db.refresh(code); return AccountCodeResponse.model_validate(code)


@router.delete("/account-codes/{code_id}", status_code=204)
async def delete_account_code(code_id: int, current_user: User = Depends(require_permission("intendance.expenses.manage_refs")), db: AsyncSession = Depends(get_db)):
    code = await db.get(IntendanceAccountCode, code_id)
    if not code: raise HTTPException(404, "Code comptable introuvable")
    await db.delete(code); await db.commit()


@router.post("/imports", response_model=ImportBatchResponse, status_code=201)
async def analyze_import(file: UploadFile = File(...), current_user: User = Depends(require_permission("intendance.expenses.import")), db: AsyncSession = Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(400, "Seuls les fichiers .xlsx sont acceptés")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(file.file, read_only=True, data_only=True)
    except Exception as exc: raise HTTPException(400, f"Fichier Excel invalide : {exc}")
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    header_aliases = {
        "Req Number": "A", "Detail Required Date": "B", "Bank Card Description": "C",
        "Item Description": "D", "Quantity": "E", "UOM Code": "F",
        "DetailExpectedPrice": "G", "Corporation": "H", "Cost Center Code": "I",
        "Cost Center": "I", "Detail Account Description": "J", "Prepared By Name": "K",
    }
    normalized_headers = {normalize_text(header).lower(): index for index, header in enumerate(headers) if header is not None}
    required_headers = ["Req Number", "Detail Required Date", "Item Description", "Quantity", "UOM Code", "DetailExpectedPrice", "Corporation", "Detail Account Description", "Prepared By Name"]
    missing_headers = [header for header in required_headers if header.lower() not in normalized_headers]
    if missing_headers or not ("Cost Center Code".lower() in normalized_headers or "Cost Center".lower() in normalized_headers):
        raise HTTPException(400, "En-têtes Excel inattendus")
    column_indexes = {letter: normalized_headers.get(header.lower()) for header, letter in header_aliases.items() if header.lower() in normalized_headers}
    batch = IntendanceImportBatch(reference=uuid.uuid4().hex, file_name=file.filename, uploaded_by=current_user.id)
    db.add(batch); await db.flush()
    seen = set(); rows = list(sheet.iter_rows(min_row=2, values_only=True)); previous_date = None
    mapped_rows = []
    for values in rows:
        mapped_values = {letter: (values[index] if index is not None and index < len(values) else None) for letter, index in column_indexes.items()}
        mapped_values.setdefault("C", None)
        mapped_rows.append(mapped_values)
    for offset, mapped_values in enumerate(mapped_rows):
        number = offset + 2
        current_date = normalize_date(mapped_values.get("B"))
        if current_date is None and previous_date is not None:
            mapped_values["B"] = previous_date
        elif current_date is None:
            for next_values in mapped_rows[offset + 1:]:
                next_date = normalize_date(next_values.get("B"))
                if next_date is not None:
                    mapped_values["B"] = next_date
                    break
        elif current_date is not None:
            previous_date = current_date
        parsed = parse_excel_row(mapped_values, number); fp = expense_fingerprint(parsed); parsed["fingerprint"] = fp
        parsed["is_duplicate_in_file"] = fp in seen; seen.add(fp)
        db.add(IntendanceImportStagedRow(batch_id=batch.id, **{k: v for k, v in parsed.items() if k in {"line_number","fingerprint","req_number","expense_date","bank_card","item_description","quantity","uom","amount","corporation","cost_center","account_code","account_label","prepared_by","invalid_reason","is_duplicate_in_file"}}))
    existing_rows = (await db.execute(select(IntendanceExpense.expense_date, IntendanceExpense.account_label, IntendanceExpense.item_description, IntendanceExpense.amount))).all()
    existing = {
        expense_duplicate_key({"expense_date": expense_date, "account_label": account_label, "item_description": item_description, "amount": amount})
        for expense_date, account_label, item_description, amount in existing_rows
    }
    staged = (await db.execute(select(IntendanceImportStagedRow).where(IntendanceImportStagedRow.batch_id == batch.id))).scalars().all()
    for row in staged:
        row.is_duplicate_existing = expense_duplicate_key(row) in existing
    batch.rows_total = len(staged); batch.rows_new = sum(not r.invalid_reason and not r.is_duplicate_existing and not r.is_duplicate_in_file for r in staged); batch.rows_duplicate_existing = sum(r.is_duplicate_existing for r in staged); batch.rows_duplicate_in_file = sum(r.is_duplicate_in_file for r in staged); batch.rows_invalid = sum(bool(r.invalid_reason) for r in staged)
    await db.commit(); await db.refresh(batch)
    return ImportBatchResponse(
        id=batch.id, reference=batch.reference, file_name=batch.file_name, status=batch.status,
        created_at=batch.created_at, uploaded_by=batch.uploaded_by,
        bilan=ImportBilan(
            rows_total=batch.rows_total, rows_new=batch.rows_new,
            rows_duplicate_existing=batch.rows_duplicate_existing,
            rows_duplicate_in_file=batch.rows_duplicate_in_file,
            rows_unknown_codes=batch.rows_unknown_codes,
            rows_unknown_users=batch.rows_unknown_users,
            rows_invalid=batch.rows_invalid,
        ),
    )


@router.post("/imports/{batch_id}/validate", response_model=ImportValidateResponse)
async def validate_import(batch_id: int, current_user: User = Depends(require_permission("intendance.expenses.import")), db: AsyncSession = Depends(get_db)):
    batch = (await db.execute(select(IntendanceImportBatch).where(IntendanceImportBatch.id == batch_id).with_for_update())).scalar_one_or_none()
    if not batch: raise HTTPException(404, "Import introuvable")
    if batch.status == "validated": return ImportValidateResponse(reference=batch.reference, created_count=batch.created_count or 0, skipped_duplicates=batch.rows_total - (batch.created_count or 0))
    rows = (await db.execute(select(IntendanceImportStagedRow).where(IntendanceImportStagedRow.batch_id == batch.id))).scalars().all(); created = 0
    aliases = {x.alias: x.user_name for x in (await db.execute(select(IntendanceCardAlias))).scalars().all()}; codes = {x.code: x for x in (await db.execute(select(IntendanceAccountCode))).scalars().all()}
    for account_code in {row.account_code for row in rows if row.account_code and row.account_code not in codes}:
        code = IntendanceAccountCode(code=account_code, category="Non affecté")
        db.add(code); codes[account_code] = code
    for row in rows:
        if row.invalid_reason or row.is_duplicate_existing or row.is_duplicate_in_file: continue
        code = codes.get(row.account_code or "")
        db.add(IntendanceExpense(fingerprint=row.fingerprint, req_number=row.req_number, expense_date=row.expense_date, bank_card=row.bank_card, item_description=row.item_description, quantity=row.quantity, uom=row.uom, amount=row.amount, corporation=row.corporation, cost_center=row.cost_center, account_code=row.account_code, account_label=row.account_label, prepared_by=row.prepared_by, category=code.category if code else None, user_name=aliases.get(row.bank_card), import_batch_id=batch.id, created_by=current_user.id)); created += 1
    batch.status = "validated"; batch.validated_at = datetime.now(timezone.utc); batch.created_count = created; await db.commit()
    return ImportValidateResponse(reference=batch.reference, created_count=created, skipped_duplicates=len(rows) - created)


@router.post("/reassign", response_model=dict)
async def reassign_references(current_user: User = Depends(require_permission("intendance.expenses.import")), db: AsyncSession = Depends(get_db)):
    """Réapplique les référentiels configurés sans écraser les corrections manuelles."""
    expenses = (await db.execute(select(IntendanceExpense).with_for_update())).scalars().all()
    aliases = {x.alias: x.user_name for x in (await db.execute(select(IntendanceCardAlias))).scalars().all()}
    codes = {x.code: x.category for x in (await db.execute(select(IntendanceAccountCode))).scalars().all()}
    missing_codes = {expense.account_code for expense in expenses if expense.account_code and expense.account_code not in codes}
    for account_code in missing_codes:
        db.add(IntendanceAccountCode(code=account_code, category="Non affecté"))
        codes[account_code] = "Non affecté"
    categories = users = 0
    for expense in expenses:
        if not expense.category_is_manual:
            category = codes.get(expense.account_code or "")
            if expense.category != category:
                expense.category = category
                categories += 1
        if not expense.user_is_manual:
            user_name = aliases.get(expense.bank_card)
            if expense.user_name != user_name:
                expense.user_name = user_name
                users += 1
    await db.commit()
    return {"categories_updated": categories, "users_updated": users, "codes_created": len(missing_codes)}

